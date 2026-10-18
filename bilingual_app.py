#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
対訳表生成ツール（Streamlit アプリ）

英語・日本語テキストファイル（<<<PAGE N>>> マーカー付き）を入力として受け取り、
センテンス単位に分割した対訳表 Excel を生成する。
"""

import io
import re

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sentence_transformers import SentenceTransformer, util

# ─── ページ設定 ────────────────────────────────────────────
st.set_page_config(
    page_title="対訳表生成ツール",
    page_icon="📄",
    layout="wide",
)


# ─── モデルキャッシュ ──────────────────────────────────────
@st.cache_resource(show_spinner="意味的照合モデルをロード中…")
def load_model():
    return SentenceTransformer("paraphrase-multilingual-MiniLM-L12-v2")


# ─── テキスト処理ユーティリティ ────────────────────────────

def parse_pages(text: str) -> dict[int, str]:
    """<<<PAGE N>>> マーカーでテキストをページ辞書に分割する。"""
    parts = re.split(r"<<<PAGE (\d+)>>>", text)
    pages = {}
    for i in range(1, len(parts), 2):
        pages[int(parts[i])] = parts[i + 1].strip()
    return pages


def extract_lines(page_content: str) -> list[str]:
    """ページテキストから意味のある行を抽出する。"""
    result = []
    for line in page_content.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("［") or line.startswith("【") or line.startswith("図"):
            continue
        if re.match(r"^\d+$", line):
            continue
        result.append(line)
    return result


def is_heading(text: str) -> bool:
    """見出し行かどうかを判定する。"""
    t = text.strip()
    if not t:
        return False
    if len(t) <= 60 and not t.endswith(".") and not t.endswith("。"):
        return True
    heading_starts = [
        "Environment", "Management", "Governance", "Strategies", "Approach",
        "Activities", "Related Information", "Initiatives", "Future Initiatives",
        "Risk Management", "Climate Change", "Sustainable", "Packaging",
        "Water", "Agricultural", "Natural Capital", "Biodiversity",
        "Communities", "Responsible Drinking", "Health", "Human Rights",
        "Enhancement", "Data Section", "GRI", "SASB",
        "環境", "マネジメント", "ガバナンス", "戦略", "考え方", "取り組み",
        "関連情報", "今後の取り組み", "リスク管理", "気候変動", "持続可能",
        "コミュニティ", "責任ある飲酒", "健康", "人権", "人的資本",
        "データセクション",
    ]
    for kw in heading_starts:
        if t.startswith(kw):
            return True
    return False


def split_sentences(text: str) -> list[str]:
    """英文をセンテンス単位に分割する。"""
    sentences = re.split(r"(?<=\.)\s+(?=[A-Z\"\'(])", text)
    return [s.strip() for s in sentences if s.strip()]


def build_jp_alignment(offset_start: int, offset_amount: int):
    """英語ページ → 日本語ページの変換関数を返す。"""
    def jp_alignment(en_page: int) -> int:
        if en_page < offset_start:
            return en_page
        return en_page - offset_amount
    return jp_alignment


# ─── 対訳表生成 ───────────────────────────────────────────

def build_bilingual_rows(
    en_pages: dict,
    jp_pages: dict,
    jp_alignment,
    model,
    progress_bar=None,
) -> list[list]:
    """
    英語・日本語ページ辞書から対訳行リストを生成する。

    Returns:
        [[seq, 種別, 英語, 日本語], ...]
    """
    rows = []
    seq = 0
    semantic_count = 0

    sorted_en_pages = sorted(en_pages.keys())
    total = len(sorted_en_pages)

    for idx, en_page_num in enumerate(sorted_en_pages):
        jp_page_num = jp_alignment(en_page_num)

        if jp_page_num not in jp_pages:
            continue

        en_lines = extract_lines(en_pages[en_page_num])
        jp_lines = extract_lines(jp_pages[jp_page_num])

        if len(en_lines) < 2:
            continue

        # JP 埋め込みをページ単位で計算
        jp_embeddings = None
        if jp_lines:
            jp_embeddings = model.encode(
                jp_lines, convert_to_tensor=True, show_progress_bar=False
            )

        for i, en_line in enumerate(en_lines):
            if i < len(jp_lines):
                jp_text = jp_lines[i]
            elif jp_embeddings is not None:
                # 意味的照合：コサイン類似度が最大の JP 行を選択
                en_emb = model.encode(
                    en_line, convert_to_tensor=True, show_progress_bar=False
                )
                sims = util.cos_sim(en_emb, jp_embeddings)[0]
                best_idx = int(sims.argmax())
                jp_text = jp_lines[best_idx]
                semantic_count += 1
            else:
                jp_text = ""

            if is_heading(en_line):
                seq += 1
                rows.append([seq, "見出し", en_line, jp_text])
            else:
                for sentence in split_sentences(en_line):
                    seq += 1
                    rows.append([seq, "段落", sentence, jp_text])

        if progress_bar is not None:
            progress_bar.progress(
                (idx + 1) / total,
                text=f"処理中… ページ {en_page_num} / {max(sorted_en_pages)}  (意味的照合: {semantic_count} 行)",
            )

    return rows


# ─── Excel 生成 ───────────────────────────────────────────

def generate_excel(rows: list[list]) -> bytes:
    """対訳行リストから Excel バイト列を生成する。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "対訳表"

    # ヘッダー
    headers = ["通し番号", "種別", "英語", "日本語"]
    ws.append(headers)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font_white = Font(bold=True, color="FFFFFF")
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # データ行
    heading_fill = PatternFill("solid", fgColor="D9E1F2")
    heading_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    for row in rows:
        ws.append(row)
        row_idx = ws.max_row
        is_h = row[1] == "見出し"
        for col in range(1, 5):
            cell = ws.cell(row=row_idx, column=col)
            cell.alignment = wrap
            if is_h:
                cell.font = heading_font
                cell.fill = heading_fill

    # 列幅
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 60

    # 先頭行固定
    ws.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ─── Streamlit UI ──────────────────────────────────────────

st.title("📄 対訳表生成ツール")
st.caption("英語・日本語テキストファイルからセンテンス単位の対訳表 Excel を生成します。")

# サイドバー：入力と設定
with st.sidebar:
    st.header("① ファイルのアップロード")
    en_file = st.file_uploader(
        "英語テキストファイル（<<<PAGE N>>> マーカー付き）",
        type=["txt"],
        key="en_file",
    )
    jp_file = st.file_uploader(
        "日本語テキストファイル（<<<PAGE N>>> マーカー付き）",
        type=["txt"],
        key="jp_file",
    )

    st.divider()
    st.header("② ページアライメント設定")
    st.caption("英語と日本語でページがずれている場合に設定します。")

    use_offset = st.checkbox("ページオフセットあり", value=False)
    offset_start = 147
    offset_amount = 1
    if use_offset:
        offset_start = st.number_input(
            "オフセット開始ページ（この英語ページ以降でずれが発生）",
            min_value=1, max_value=500, value=147, step=1,
        )
        offset_amount = st.number_input(
            "オフセット量（日本語が何ページ前にあるか）",
            min_value=1, max_value=10, value=1, step=1,
        )
        st.info(
            f"英語 p{offset_start} 以降：JP = EN − {offset_amount}\n\n"
            f"例）英語 p{offset_start} → 日本語 p{offset_start - offset_amount}"
        )

    st.divider()
    run_button = st.button(
        "▶ 対訳表を生成",
        type="primary",
        disabled=(en_file is None or jp_file is None),
        use_container_width=True,
    )

# メインエリア
if en_file is None or jp_file is None:
    st.info("サイドバーから英語・日本語テキストファイルをアップロードしてください。")
    st.markdown("""
    ### ファイル形式
    テキストファイルに `<<<PAGE N>>>` マーカーを含めてください。

    ```
    <<<PAGE 1>>>
    Introduction
    This report describes our sustainability initiatives...

    <<<PAGE 2>>>
    Environment
    We reduced CO2 emissions by 15% compared to the previous year.
    ...
    ```
    """)
    st.stop()

if run_button:
    # テキスト読み込み
    try:
        en_text = en_file.read().decode("utf-8")
        jp_text = jp_file.read().decode("utf-8")
    except UnicodeDecodeError:
        st.error("ファイルの文字コードが UTF-8 ではありません。UTF-8 で保存し直してください。")
        st.stop()

    en_pages = parse_pages(en_text)
    jp_pages = parse_pages(jp_text)

    if not en_pages:
        st.error("英語ファイルに <<<PAGE N>>> マーカーが見つかりません。")
        st.stop()
    if not jp_pages:
        st.error("日本語ファイルに <<<PAGE N>>> マーカーが見つかりません。")
        st.stop()

    st.success(
        f"ファイル読み込み完了：英語 {len(en_pages)} ページ / 日本語 {len(jp_pages)} ページ"
    )

    # モデルロード
    model = load_model()

    # 対訳表生成
    jp_align_func = build_jp_alignment(offset_start, offset_amount)
    progress_bar = st.progress(0, text="処理を開始しています…")

    with st.spinner("対訳表を生成中…"):
        rows = build_bilingual_rows(en_pages, jp_pages, jp_align_func, model, progress_bar)

    progress_bar.progress(1.0, text="完了！")

    # 統計
    total = len(rows)
    with_jp = sum(1 for r in rows if r[3].strip())
    headings = sum(1 for r in rows if r[1] == "見出し")
    paragraphs = total - headings

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("総行数", f"{total:,}")
    col2.metric("日本語カバレッジ", f"{with_jp * 100 // total}%")
    col3.metric("見出し行", f"{headings:,}")
    col4.metric("段落行", f"{paragraphs:,}")

    # プレビュー
    st.subheader("プレビュー（先頭 100 行）")
    df = pd.DataFrame(rows[:100], columns=["通し番号", "種別", "英語", "日本語"])
    st.dataframe(df, use_container_width=True, height=400)

    # Excel ダウンロード
    st.subheader("ダウンロード")
    excel_bytes = generate_excel(rows)
    st.download_button(
        label="📥 Excel ファイルをダウンロード",
        data=excel_bytes,
        file_name="bilingual_table.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )

elif en_file is not None and jp_file is not None:
    # ファイルがアップロードされたが未実行の状態
    en_text = en_file.read().decode("utf-8", errors="replace")
    jp_text = jp_file.read().decode("utf-8", errors="replace")
    en_pages = parse_pages(en_text)
    jp_pages = parse_pages(jp_text)
    st.info(
        f"ファイル確認済み：英語 {len(en_pages)} ページ / 日本語 {len(jp_pages)} ページ\n\n"
        "サイドバーの「▶ 対訳表を生成」ボタンを押してください。"
    )
